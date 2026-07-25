"""
연도별 지역축제 개최계획 엑셀은 시트명/컬럼 배치가 매년 다릅니다.
그래서 연도별로 별도 파서 함수를 두고, excel_loader.py에서 plan_year 기준으로 라우팅합니다.

새 연도 파일이 추가되면:
1. 이 파일에 parse_YYYY(path, plan_year) 함수를 새로 추가
2. excel_loader.py의 SCHEMA_PARSERS 딕셔너리에 등록
"""

import re
import logging
from datetime import date

import openpyxl

from utils import parse_date, to_num, to_int

log = logging.getLogger(__name__)

FIELD_KEYS = [
    "sido_name", "sigungu_name", "festival_name", "festival_type",
    "event_place_name", "event_place_type", "place_sido", "place_sigungu",
    "place_eupmyeondong", "start_date", "end_date", "period_raw_text",
    "total_days", "hold_cycle", "hold_method", "first_held_year",
    "budget_total_mil", "budget_gov_mil", "budget_local_mil", "budget_etc_mil",
    "gov_support_dept", "visitor_total", "visitor_domestic", "visitor_foreign",
    "visitor_measure_method", "org_name", "org_type", "manager_dept",
    "manager_name", "manager_contact", "note",
]


def _empty_row(row_no):
    return {"row_no": row_no, **{k: None for k in FIELD_KEYS}}


def parse_period_text(text, plan_year):
    if not text:
        return None, None
    text = str(text).strip()


    m = re.match(r"^(\d{1,2})\s*\.\s*(\d{1,2})\s*\.?\s*~\s*(\d{1,2})\s*\.\s*(\d{1,2})\s*\.?$", text)
    if m:
        sm, sd, em, ed = (int(g) for g in m.groups())
        try:
            start = date(plan_year, sm, sd)
            end_year = plan_year if em >= sm else plan_year + 1
            end = date(end_year, em, ed)
            return start, end
        except ValueError:
            return None, None


    m = re.match(r"^(\d{1,2})\s*\.\s*(\d{1,2})\s*\.?\s*~\s*(\d{1,2})$", text)
    if m:
        sm, sd, ed = (int(g) for g in m.groups())
        try:
            start = date(plan_year, sm, sd)
            end = date(plan_year, sm, ed)
            return start, end
        except ValueError:
            return None, None


    m = re.match(r"^(\d{1,2})\s*\.\s*(\d{1,2})\s*\.?$", text)
    if m:
        mo, d = (int(g) for g in m.groups())
        try:
            single = date(plan_year, mo, d)
            return single, single
        except ValueError:
            return None, None

    return None, None 


def parse_2026(path, plan_year):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["조사표"]
    rows_out = []

    for row in ws.iter_rows(min_row=8, values_only=True):
        festival_name = row[4]
        if not festival_name:
            continue

        r = _empty_row(row[1])
        r.update({
            "sido_name": row[2], "sigungu_name": row[3],
            "festival_name": festival_name, "festival_type": row[5],
            "event_place_name": row[6], "event_place_type": row[7],
            "place_sido": row[8], "place_sigungu": row[9], "place_eupmyeondong": row[10],
            "start_date": parse_date(row[11], row[12], row[13]),
            "end_date": parse_date(row[14], row[15], row[16]),
            "total_days": to_int(row[17]), "hold_cycle": row[19],
            "first_held_year": to_int(row[20]),
            "budget_total_mil": to_num(row[21]), "budget_gov_mil": to_num(row[22]),
            "budget_local_mil": to_num(row[23]), "budget_etc_mil": to_num(row[24]),
            "gov_support_dept": row[25],
            "visitor_total": to_num(row[26]), "visitor_domestic": to_num(row[27]),
            "visitor_foreign": to_num(row[28]), "visitor_measure_method": row[29],
            "org_name": row[31], "org_type": row[32],
            "manager_dept": row[35], "manager_name": row[37], "manager_contact": row[38],
            "note": row[39],
        })
        rows_out.append(r)

    return rows_out



def parse_2025(path, plan_year):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["조사표"]
    rows_out = []

    for row in ws.iter_rows(min_row=8, values_only=True):
        festival_name = row[4]
        if not festival_name:
            continue

        r = _empty_row(row[1])
        r.update({
            "sido_name": row[2], "sigungu_name": row[3],
            "festival_name": festival_name, "festival_type": row[5],
            "event_place_name": row[6], "event_place_type": row[7],
            "place_sido": row[8], "place_sigungu": row[9], "place_eupmyeondong": row[10],
            "start_date": parse_date(row[11], row[12], row[13]),
            "end_date": parse_date(row[14], row[15], row[16]),
            "total_days": to_int(row[17]), "hold_cycle": row[19], "hold_method": row[20],
            "first_held_year": to_int(row[21]),
            "budget_total_mil": to_num(row[22]), "budget_gov_mil": to_num(row[23]),
            "budget_local_mil": to_num(row[24]), "budget_etc_mil": to_num(row[25]),
            "visitor_total": to_num(row[26]), "visitor_domestic": to_num(row[27]),
            "visitor_foreign": to_num(row[28]),
            "org_name": row[29], "org_type": row[30],
            "gov_support_dept": row[32],
            "manager_dept": row[34], "manager_contact": row[35],
        })
        rows_out.append(r)

    return rows_out


def parse_2024(path, plan_year):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["세부현황"]
    rows_out = []

    for row in ws.iter_rows(min_row=7, values_only=True):
        festival_name = row[4]
        if not festival_name:
            continue

        period_text = row[6]
        start_date, end_date = parse_period_text(period_text, plan_year)

        r = _empty_row(row[1])
        r.update({
            "sido_name": row[2], "sigungu_name": row[3],
            "festival_name": festival_name, "festival_type": row[5],
            "event_place_name": row[7],
            "start_date": start_date, "end_date": end_date,
            "period_raw_text": period_text if not start_date else None,
            "hold_method": row[8],
            "first_held_year": to_int(row[9]),
            "budget_total_mil": to_num(row[10]), "budget_gov_mil": to_num(row[11]),
            "budget_local_mil": to_num(row[12]), "budget_etc_mil": to_num(row[13]),
            "visitor_total": to_num(row[14]), "visitor_domestic": to_num(row[15]),
            "visitor_foreign": to_num(row[16]),
            "org_name": row[17], "org_type": row[18],
            "gov_support_dept": row[20],
            "manager_dept": row[22], "manager_name": row[24], "manager_contact": row[25],
        })
        rows_out.append(r)

    return rows_out
